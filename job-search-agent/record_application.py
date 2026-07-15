#!/usr/bin/env python3
"""record_application.py — Registra una postulación en ambos trackers.

Agrega una entrada a job_tracker.json Y una fila a job_tracker.xlsx (hoja
"Postulaciones"), manteniéndolos sincronizados. Antes de registrar:

  1. Ejecuta la lógica anti-duplicado (misma que check_company.py):
     empresa excluida / misma empresa dentro de la ventana / jd_url repetida.
  2. Respeta el cap diario y el max_applications por search_query
     (las entradas con status "Incompleto..." no cuentan).

Uso típico (postulación exitosa):
    python record_application.py \
        --company "Acme Corp" \
        --role "AI Engineer" \
        --role-profile ai_engineer \
        --search-query "AI Engineer" \
        --location "Remoto - Global" \
        --platform linkedin \
        --jd-url "https://www.linkedin.com/jobs/view/123" \
        --status "Aplicado" \
        --notes "Easy Apply"

Registrar una incompleta (no cuenta para el cap):
    python record_application.py ... --status "Incompleto - CAPTCHA"

Flags útiles:
    --date YYYY-MM-DD   fecha de la postulación (default: hoy)
    --force             saltea las verificaciones anti-duplicado/cap
    --dry-run           muestra qué haría sin escribir nada
"""

import argparse
import json
import os
import sys
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
TRACKER_JSON = os.path.join(HERE, "job_tracker.json")
TRACKER_XLSX = os.path.join(HERE, "job_tracker.xlsx")
CONFIG_PATH = os.path.join(HERE, "config.json")

XLSX_COLUMNS = [
    "id", "date", "company", "role", "role_profile", "search_query",
    "location", "platform", "jd_url", "resume_pdf", "status", "notes",
]

RESUME_LABELS = {
    "ai_engineer": "AI-Engineer",
    "automation_engineer": "Automation-Engineer",
    "backend_python": "Python-Backend",
    "fullstack": "Full-Stack",
}


def _norm(v):
    return (v or "").strip().lower()


def _load_json(path, default):
    if not os.path.exists(path):
        return default
    with open(path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def _parse_date(v):
    try:
        return datetime.strptime((v or "")[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None


def _is_valid_status(status):
    """Una postulación 'válida' (cuenta para el cap) no es Incompleto."""
    return not _norm(status).startswith("incompleto")


def _next_id(tracker, today_str):
    todays = [e for e in tracker if e.get("date") == today_str]
    return f"{today_str}-{len(todays) + 1:03d}"


def _default_resume(role_profile):
    label = RESUME_LABELS.get(role_profile, role_profile)
    return f"tailored_resumes/JoseDaLuzPereira_{label}.pdf"


def check_duplicates(tracker, config, company, role_profile, jd_url, days):
    company_n = _norm(company)

    excluded = [_norm(c) for c in config.get("excluded_companies", [])]
    if company_n in excluded:
        return f"Regla 4: '{company}' está en excluded_companies."

    if jd_url:
        for e in tracker:
            if _norm(e.get("jd_url")) == _norm(jd_url):
                return f"Regla 3: jd_url ya registrada (id={e.get('id')})."

    today = date.today()
    for e in tracker:
        if _norm(e.get("company")) != company_n:
            continue
        d = _parse_date(e.get("date"))
        if d is None:
            continue
        ago = (today - d).days
        if ago < 0 or ago > days:
            continue
        same = e.get("role_profile") == role_profile
        rule = "Regla 1 (mismo perfil)" if same else "Regla 2 (distinto perfil)"
        return f"{rule}: postulación a '{company}' hace {ago} día(s) (id={e.get('id')})."
    return None


def check_caps(tracker, config, search_query, today_str):
    """Devuelve (error|None)."""
    cap = config.get("daily_application_cap", 10)
    valid_today = [e for e in tracker
                   if e.get("date") == today_str and _is_valid_status(e.get("status"))]
    if len(valid_today) >= cap:
        return f"cap diario alcanzado ({len(valid_today)}/{cap})."

    run = next((r for r in config.get("search_runs", [])
                if r.get("search_query") == search_query), None)
    if run:
        maxq = run.get("max_applications", cap)
        applied_q = [e for e in valid_today if e.get("search_query") == search_query]
        if len(applied_q) >= maxq:
            return (f"query '{search_query}' agotó su cupo "
                    f"({len(applied_q)}/{maxq}).")
    return None


def append_xlsx(entry):
    from openpyxl import load_workbook
    wb = load_workbook(TRACKER_XLSX)
    ws = wb["Postulaciones"] if "Postulaciones" in wb.sheetnames else wb.active
    ws.append([entry.get(col, "") for col in XLSX_COLUMNS])
    wb.save(TRACKER_XLSX)


def main():
    p = argparse.ArgumentParser(description="Registra una postulación en los trackers.")
    p.add_argument("--company", required=True)
    p.add_argument("--role", required=True)
    p.add_argument("--role-profile", required=True)
    p.add_argument("--search-query", required=True)
    p.add_argument("--location", default="")
    p.add_argument("--platform", required=True, choices=["linkedin", "indeed"])
    p.add_argument("--jd-url", required=True)
    p.add_argument("--resume-pdf", default=None)
    p.add_argument("--status", default="Aplicado")
    p.add_argument("--notes", default="")
    p.add_argument("--date", default=None, help="YYYY-MM-DD (default: hoy)")
    p.add_argument("--force", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    config = _load_json(CONFIG_PATH, {})
    tracker = _load_json(TRACKER_JSON, [])
    today_str = args.date or date.today().isoformat()
    days = config.get("duplicate_check_days_threshold", 60)

    counts_for_cap = _is_valid_status(args.status)

    if not args.force:
        dup = check_duplicates(tracker, config, args.company,
                               args.role_profile, args.jd_url, days)
        if dup:
            print(f"RECHAZADO (anti-duplicado) — {dup}")
            return 1
        # Las incompletas no consumen cupo, así que no se validan contra el cap.
        if counts_for_cap:
            cap_err = check_caps(tracker, config, args.search_query, today_str)
            if cap_err:
                print(f"RECHAZADO (cap) — {cap_err}")
                return 1

    entry = {
        "id": _next_id(tracker, today_str),
        "date": today_str,
        "company": args.company,
        "role": args.role,
        "role_profile": args.role_profile,
        "search_query": args.search_query,
        "location": args.location,
        "platform": args.platform,
        "jd_url": args.jd_url,
        "resume_pdf": args.resume_pdf or _default_resume(args.role_profile),
        "status": args.status,
        "notes": args.notes,
    }

    if args.dry_run:
        print("DRY-RUN — se registraría:")
        print(json.dumps(entry, ensure_ascii=False, indent=2))
        return 0

    tracker.append(entry)
    with open(TRACKER_JSON, "w", encoding="utf-8") as fh:
        json.dump(tracker, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    append_xlsx(entry)

    print(f"REGISTRADO — id={entry['id']} | {entry['company']} | "
          f"{entry['status']} | {entry['resume_pdf']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
